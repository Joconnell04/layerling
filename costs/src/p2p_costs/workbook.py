"""Build and read costs/cost_model.xlsx.

Ownership:
- Inputs, Market: Jackson's. Built once; refresh only appends missing input keys and new research
  listings (and rewrites Market's formula columns I and T onward).
- Prices column N (Override): Jackson's. Refresh keeps each value by the row key in column U.
- Test data Use column: Jackson's override per run (1 in, 0 out, blank = Auto). Refresh never
  writes to an existing Use cell; it updates the other columns in place and appends new runs.
- Per product, Prices: rebuilt by refresh with live formulas. Test data: updated in place.
- Changelog: append only.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .defaults import (
    CHANNEL_LABELS,
    CHANNELS,
    JACKSON,
    M_TIER,
    MARKET_FOR,
    MARKET_PRODUCTS,
    PRODUCT_LABELS,
    PRODUCTS,
    SIZE_PRODUCTS,
    SIZE_TIERS_MM,
    TIERS,
    Input,
)
from .ingest import RunRow
from .market import default_use

INPUTS, TEST, MARKET, PER, PRICES, SIZES, LOG = (
    "Inputs", "Test data", "Market", "Per product", "Prices", "Sizes", "Changelog")
SHEET_ORDER = [INPUTS, TEST, MARKET, PER, PRICES, SIZES, LOG]
HDR = 4  # header row on Inputs, Test data, Market
FIRST = 5  # first data row
TD_LAST = 5000
MK_LAST = 2000

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="666666")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_KEY = Font(name=FONT, size=9, color="808080")
F_URL = Font(name=FONT, size=9, color="0563C1", underline="single")
FILL_HDR = PatternFill("solid", fgColor="333F4F")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
FILL_RESULT = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

USD = '$#,##0.00;($#,##0.00);"-"'
USD3 = '$#,##0.000;($#,##0.000);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'


def fmt_for_unit(unit: str) -> str:
    u = unit.lower()
    if u == "share":
        return PCT
    if u.startswith("usd/credit"):
        return "$0.0000"
    if u.startswith("usd/kwh"):
        return "$0.000"
    if u.startswith("usd"):
        return USD
    if u in ("g", "oz", "multiplier"):
        return "0.0#"
    return "General"


def q(sheet: str) -> str:
    return f"'{sheet}'" if " " in sheet else sheet


def _header(ws: Worksheet, row: int, labels: list[str], widths: list[float] | None = None) -> None:
    for j, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _title(ws: Worksheet, title: str, legend: str) -> None:
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = legend
    ws["A2"].font = F_NOTE


# ------------------------------------------------------------------ Inputs

INPUT_COLS = ["Section", "Input", "Value", "Unit", "Provenance", "Source", "Notes",
              "If wrong, what moves", "Key"]
# Column number of each Inputs field. Formulas point at the Value column.
IC = {"section": 1, "label": 2, "value": 3, "unit": 4, "provenance": 5, "source": 6, "notes": 7,
      "moves": 8, "key": 9}
VALUE_COL = get_column_letter(IC["value"])


def write_input_row(ws: Worksheet, r: int, inp: Input) -> None:
    vals = {"section": inp.section, "label": inp.label, "value": inp.value, "unit": inp.unit,
            "provenance": inp.provenance, "source": inp.source, "notes": inp.notes,
            "moves": inp.moves, "key": inp.key}
    for k, v in vals.items():
        c = ws.cell(row=r, column=IC[k], value=v)
        c.font = F_BASE
        c.border = BOX
        c.alignment = WRAP
    ws.cell(row=r, column=IC["key"]).font = F_KEY
    v = ws.cell(row=r, column=IC["value"])
    v.font = F_INPUT
    v.number_format = fmt_for_unit(inp.unit)
    if inp.unit == "share" and isinstance(inp.value, float) and round(inp.value * 1000, 6) % 1:
        v.number_format = "0.000%"
    if inp.provenance == JACKSON:
        v.fill = FILL_YELLOW
    src = ws.cell(row=r, column=IC["source"])
    if inp.source.startswith("http"):
        src.hyperlink = inp.source
        src.font = F_URL


def build_inputs(ws: Worksheet, inputs: list[Input]) -> None:
    _title(ws, "Inputs",
           "Edit the Value column (blue). Yellow cells are yours to fill. Provenance: SOURCE (with "
           "link), FROM JACKSON, or ASSUMED. Refresh never changes this sheet; it only adds new "
           "keys at the bottom.")
    _header(ws, HDR, INPUT_COLS, [18, 46, 12, 13, 14, 36, 60, 44, 30])
    ws.freeze_panes = ws.cell(row=FIRST, column=IC["value"])
    for n, inp in enumerate(inputs):
        write_input_row(ws, FIRST + n, inp)


def read_inputs(ws_f: Worksheet, ws_v: Worksheet | None = None) -> dict[str, dict]:
    """key -> {row, value, provenance, source, ...}. Values come from the cached-values sheet when
    given (so a formula Jackson typed into a Value cell reads as its result)."""
    out: dict[str, dict] = {}
    for r in range(FIRST, ws_f.max_row + 1):
        key = ws_f.cell(row=r, column=IC["key"]).value
        if not key:
            continue
        raw = ws_f.cell(row=r, column=IC["value"]).value
        val = ws_v.cell(row=r, column=IC["value"]).value if ws_v is not None else raw
        if isinstance(raw, str) and raw.startswith("=") and ws_v is None:
            val = None
        d = {k: ws_f.cell(row=r, column=c).value for k, c in IC.items() if k != "value"}
        d.update(row=r, value=val, raw=raw)
        out[str(key).strip()] = d
    return out


def append_missing_inputs(ws: Worksheet, defaults: list[Input], have: dict[str, dict]) -> list[str]:
    added = []
    r = max([d["row"] for d in have.values()] + [HDR]) + 1
    for inp in defaults:
        if inp.key not in have:
            write_input_row(ws, r, inp)
            ws.cell(row=r, column=IC["notes"]).value = f"{inp.notes} (Added by refresh.)".strip()
            have[inp.key] = {"row": r, "value": inp.value, "provenance": inp.provenance}
            added.append(inp.key)
            r += 1
    return added


# ------------------------------------------------------------------ Test data

TD_COLS = [
    ("run_key", "Run key", 30), ("product", "Product", 13), ("kind", "Kind", 11),
    ("config", "Config (settings, variant, orientation)", 48),
    ("use", "Use (yours: 1 in, 0 out, blank auto)", 12), ("auto", "Auto", 7),
    ("counts", "Counts", 8), ("credits", "Credits", 9), ("meshy_usd", "Meshy USD (ledger)", 11),
    ("pla_bambu_g", "Bambu PLA g", 10), ("pla_generic_g", "Generic PLA g", 10),
    ("petg_g", "PETG g", 9), ("other_g", "Other g", 9), ("total_g", "Total g (incl. purge)", 11),
    ("purge_g", "Purge g", 9), ("print_min", "Print min", 9), ("slots", "Slots", 7),
    ("ok", "Slice ok", 8), ("filaments", "Grams per filament", 44), ("source", "Source file", 44),
    ("notes", "Notes", 50), ("run_at", "Run at", 20), ("tier", "Size tier", 8),
    ("size_mm", "Made at mm", 9),
]
TD_LETTER = {k: get_column_letter(j) for j, (k, _, _) in enumerate(TD_COLS, start=1)}
TD_COL = {k: j for j, (k, _, _) in enumerate(TD_COLS, start=1)}
TD_LEGEND = ("One row per run, from test outputs. Auto = 1 when the run's config matches what "
             "production does now. Use is yours: 1 forces a run in, 0 forces it out, blank "
             "follows Auto; refresh never writes to it. Counts = Use if set, else Auto. Grams per "
             "filament include purge.")


def _td_layout_ok(ws: Worksheet) -> bool:
    """The Use-override layout, possibly missing columns added later at the end (refresh adds
    them in place)."""
    have = [ws.cell(row=HDR, column=j).value for j in range(1, len(TD_COLS) + 1)]
    want = [h for _, h, _ in TD_COLS]
    n = 0
    while n < len(have) and have[n] is not None:
        n += 1
    return n > TD_COL["use"] and have[:n] == want[:n]


def _write_td_row(ws: Worksheet, r: int, d: dict, write_use: bool) -> None:
    """Every column of one run except Use, which is written only when the row is new."""
    for k, j in TD_COL.items():
        if k == "use" and not write_use:
            continue
        if k == "counts":
            u, a = TD_LETTER["use"], TD_LETTER["auto"]
            v = f"=IF(ISNUMBER({u}{r}),{u}{r},{a}{r})"
        else:
            v = d.get(k)
        if isinstance(v, bool):
            v = "yes" if v else "no"
        c = ws.cell(row=r, column=j, value=v)
        c.font = F_BASE
        c.border = BOX
        if k in ("credits", "slots"):
            c.number_format = NUM0
        elif k.endswith("_g") or k == "print_min":
            c.number_format = "0.00"
        elif k == "meshy_usd":
            c.number_format = USD
        elif k == "counts":
            c.font = F_BOLD
    u = ws.cell(row=r, column=TD_COL["use"])
    u.font = F_INPUT
    u.fill = FILL_YELLOW
    u.border = BOX


def update_test_data(ws: Worksheet, rows: list[RunRow],
                     migrated_uses: dict[str, object] | None = None) -> None:
    """Update Test data in place. Existing rows keep their Use cell untouched; new runs are
    appended with Use blank. A sheet in an older layout is rebuilt once, keeping only the Use
    values Jackson had edited (migrated_uses)."""
    fresh = not _td_layout_ok(ws)
    if fresh:
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
        ws.delete_rows(1, ws.max_row)
    _title(ws, "Test data", TD_LEGEND)
    _header(ws, HDR, [h for _, h, _ in TD_COLS], [w for _, _, w in TD_COLS])
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    where = {}
    for r in range(FIRST, ws.max_row + 1):
        k = ws.cell(row=r, column=TD_COL["run_key"]).value
        if k:
            where[str(k)] = r
    order = {p: n for n, p in enumerate(PRODUCTS)}
    nxt = max([*where.values(), HDR]) + 1
    for row in sorted(rows, key=lambda x: (order.get(x.product, 99), x.kind, x.run_key)):
        d = row.to_json()
        if row.run_key in where:
            _write_td_row(ws, where[row.run_key], d, write_use=False)
        else:
            d["use"] = (migrated_uses or {}).get(row.run_key) if fresh else None
            _write_td_row(ws, nxt, d, write_use=True)
            where[row.run_key] = nxt
            nxt += 1


def read_test_uses(ws: Worksheet) -> dict[str, object]:
    """run key -> Use cell, in either layout (the old one had Use in column D)."""
    col = TD_COL["use"] if _td_layout_ok(ws) else 4
    out = {}
    for r in range(FIRST, ws.max_row + 1):
        k = ws.cell(row=r, column=1).value
        if k:
            out[str(k)] = ws.cell(row=r, column=col).value
    return out


def read_test_rows(ws: Worksheet) -> list[dict]:
    """Rows with counts resolved from Use and Auto, as the Counts formula does."""
    out = []
    for r in range(FIRST, ws.max_row + 1):
        if not ws.cell(row=r, column=1).value:
            continue
        d = {k: ws.cell(row=r, column=j).value for j, (k, _, _) in enumerate(TD_COLS, 1)}
        d["counts"] = effective_use(d.get("use"), d.get("auto"))
        out.append(d)
    return out


def effective_use(use, auto):
    if isinstance(use, (int, float)) and not isinstance(use, bool):
        return use
    return auto


# ------------------------------------------------------------------ Market

MK_COLS = [
    ("product", "Product", 11), ("channel", "Channel", 11), ("region", "Region", 9),
    ("shop", "Shop", 22), ("shop_location", "Shop location", 16), ("title", "Title", 34),
    ("price_usd", "Price", 9), ("shipping_usd", "Shipping", 9),
    ("delivered", "Delivered price", 10), ("turnaround_days", "Turnaround days", 10),
    ("method", "Method", 16), ("size", "Size", 10), ("status", "Status", 8), ("use", "Use", 6),
    ("url", "Link", 40), ("evidence", "Evidence", 14), ("notes", "Notes", 40),
    ("accessed", "Accessed", 11),
]
MK_LETTER = {k: get_column_letter(j) for j, (k, _, _) in enumerate(MK_COLS, start=1)}
MK_ALL0 = len(MK_COLS) + 2  # first helper column (all regions), after one blank column
MK_NY0 = MK_ALL0 + len(MARKET_PRODUCTS) + 1


def mk_helper(product: str, ny: bool) -> str:
    base = MK_NY0 if ny else MK_ALL0
    return get_column_letter(base + MARKET_PRODUCTS.index(product))


def _write_market_row(ws: Worksheet, r: int, d: dict) -> None:
    d = dict(d)
    if d.get("use") in (None, ""):
        d["use"], why = default_use(d)
        if why:
            d["notes"] = f"Use 0: {why}. {d.get('notes') or ''}".strip()
    for j, (k, _, _) in enumerate(MK_COLS, start=1):
        if k == "delivered":
            continue
        v = d.get(k)
        c = ws.cell(row=r, column=j, value=v)
        c.font = F_BASE
        c.border = BOX
        if k in ("price_usd", "shipping_usd"):
            c.number_format = USD
            c.font = F_INPUT
        if k == "use":
            c.font = F_INPUT
            c.fill = FILL_YELLOW
        if k == "url" and isinstance(v, str) and v.startswith("http"):
            c.hyperlink = v
            c.font = F_URL


def build_market(ws: Worksheet, listings: list[dict]) -> None:
    _title(ws, "Market",
           "Comparable listings. Yours to edit: add rows, fix prices, set Use to 0 to drop one. "
           "Delivered price = price plus shipping (shipping blank counts as 0). Refresh rewrites "
           "column I and the helper columns on the right, and adds new research listings.")
    _header(ws, HDR, [h for _, h, _ in MK_COLS], [w for _, _, w in MK_COLS])
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    for n, d in enumerate(listings):
        _write_market_row(ws, FIRST + n, d)


def read_market(ws: Worksheet) -> list[dict]:
    out = []
    for r in range(FIRST, ws.max_row + 1):
        d = {k: ws.cell(row=r, column=j).value for j, (k, _, _) in enumerate(MK_COLS, 1)}
        if d.get("product") or d.get("url"):
            d["row"] = r
            out.append(d)
    return out


def append_market(ws: Worksheet, have: list[dict], research: list[dict]) -> int:
    urls = {str(d.get("url") or "").strip() for d in have}
    r = max([d["row"] for d in have] + [HDR]) + 1
    n = 0
    for d in research:
        u = str(d.get("url") or "").strip()
        if not u or u in urls or d.get("product") not in MARKET_PRODUCTS:
            continue
        _write_market_row(ws, r, d)
        urls.add(u)
        r += 1
        n += 1
    return n


def write_market_formulas(ws: Worksheet) -> None:
    """Delivered price (column I) and helper columns for every listing row."""
    L = MK_LETTER
    for j, p in enumerate(MARKET_PRODUCTS):
        for ny in (False, True):
            col = mk_helper(p, ny)
            c = ws[f"{col}{HDR}"]
            c.value = p
            c.font = F_HDR
            c.fill = FILL_HDR
            ws.column_dimensions[col].width = 10
    ws[f"{mk_helper(MARKET_PRODUCTS[0], False)}3"] = "Delivered price, all regions (formula)"
    ws[f"{mk_helper(MARKET_PRODUCTS[0], True)}3"] = "Delivered price, NY shops (formula)"
    for col in (mk_helper(MARKET_PRODUCTS[0], False), mk_helper(MARKET_PRODUCTS[0], True)):
        ws[f"{col}3"].font = F_BOLD
    last = max(ws.max_row, FIRST)
    for r in range(FIRST, last + 1):
        if not (ws[f"A{r}"].value or ws[f"{L['url']}{r}"].value):
            continue
        ws[f"{L['delivered']}{r}"] = (
            f'=IF(ISNUMBER({L["price_usd"]}{r}),{L["price_usd"]}{r}'
            f'+IF(ISNUMBER({L["shipping_usd"]}{r}),{L["shipping_usd"]}{r},0),"")')
        ws[f"{L['delivered']}{r}"].number_format = USD
        ws[f"{L['delivered']}{r}"].font = F_BASE
        for p in MARKET_PRODUCTS:
            a, n = mk_helper(p, False), mk_helper(p, True)
            ws[f"{a}{r}"] = (f'=IF(AND($A{r}={a}${HDR},$N{r}=1,ISNUMBER($I{r})),$I{r},"")')
            ws[f"{n}{r}"] = (f'=IF(AND($A{r}={n}${HDR},$N{r}=1,$C{r}="NY",ISNUMBER($I{r})),'
                             f'$I{r},"")')
            for col in (a, n):
                ws[f"{col}{r}"].number_format = USD
                ws[f"{col}{r}"].font = F_BASE


# ------------------------------------------------------------------ Per product


@dataclass
class Layout:
    inputs: dict[str, dict]
    per_rows: dict[str, int]
    rate_rows: dict[str, int]
    price_rows: dict[tuple[str, str], int]
    channel_rows: dict[str, int]
    market_day_rows: dict[str, int]
    size_rows: dict[tuple[str, str], int] = field(default_factory=dict)
    other_rows: dict[str, int] = field(default_factory=dict)


PER_KEYS_ROW = 15
PER_COL0 = 3  # column C holds the first product


def per_col(p: str) -> str:
    return get_column_letter(PER_COL0 + PRODUCTS.index(p))


RATES = [
    ("usd_credit", "Meshy USD per credit", "USD/credit",
     "Top-up rate if the flag is 1, else plan price / plan credits"),
    ("pla_bambu_g", "Bambu PLA USD per g", "USD/g", "Price per kg / 1000"),
    ("pla_generic_g", "Generic PLA USD per g", "USD/g", "Price per kg / 1000"),
    ("petg_g", "PETG USD per g", "USD/g", "Price per kg / 1000"),
    ("power_h", "Electricity per print hour", "USD/hour", "Power (W) / 1000 x USD per kWh"),
    ("dep_h", "Depreciation and upkeep per print hour", "USD/hour",
     "Printer price / life hours + upkeep per hour"),
    ("labor_h", "Labor rate in use", "USD/hour", "Jackson's rate if filled in, else default"),
    ("tax", "Sales tax buyers pay", "share", "NYC combined rate"),
]


def _in(lay: Layout, key: str) -> str:
    return f"{INPUTS}!${VALUE_COL}${lay.inputs[key]['row']}"


def _td(col_key: str) -> str:
    L = TD_LETTER[col_key]
    return f"{q(TEST)}!${L}${FIRST}:${L}${TD_LAST}"


def _avg_f(x: str, col_key: str, need_key: str, fallback: str, m_only: bool = True) -> str:
    tier = f"{_td('tier')},\"m\"," if m_only else ""
    crit = (f"{_td('product')},{x}${PER_KEYS_ROW},{_td('counts')},1,{tier}"
            f"{_td(need_key)},\"<>\"")
    return (f"=IF(COUNTIFS({crit})>0,SUMIFS({_td(col_key)},{crit})/COUNTIFS({crit}),"
            f"{fallback})")


def _round_f(x: str, ending: str) -> str:
    return f"ROUND(ROUNDUP(ROUND({x}-{ending},2),0)+{ending},2)"


def per_lines(lay: Layout):
    """(id, label, unit, formula(x, p) or None for a section row, how)."""
    i = lambda k: _in(lay, k)
    R = lambda k: f"$C${lay.rate_rows[k]}"
    L = lambda k, x: f"{x}{lay.per_rows[k]}"
    fail = f"(1+{i('failed_print_pct')})"
    waste = f"(1+{i('material_waste_pct')})"
    ch_own = lay.channel_rows["own_site"]
    P = q(PRICES)

    def own_row(p):
        return lay.price_rows[(p, "own_site")]

    def mk(p, ny=False):
        col = mk_helper(MARKET_FOR[p], ny)
        return f"{MARKET}!${col}${FIRST}:${col}${MK_LAST}"

    return [
        ("_s_meshy", "Meshy", None, None, ""),
        ("runs_credits", "Test runs with credits (Counts = 1)", "runs",
         lambda x, p: (f"=COUNTIFS({_td('product')},{x}${PER_KEYS_ROW},{_td('counts')},1,"
                       f"{_td('credits')},\"<>\")"), "Count in Test data"),
        ("credits", "Credits per order", "credits",
         lambda x, p: _avg_f(x, "credits", "credits", i(f"{p}.credits"), m_only=False),
         "Average of test runs, else the placeholder in Inputs"),
        ("credits_basis", "Credits from", "",
         lambda x, p: f'=IF({L("runs_credits", x)}>0,"test data","placeholder")', ""),
        ("meshy_usd", "Meshy", "USD",
         lambda x, p: f"={L('credits', x)}*(1+{i('meshy_rework_pct')})*{R('usd_credit')}",
         "Credits x (1 + rework) x USD per credit"),
        ("_s_claude", "Claude and SAM 3", None, None, ""),
        ("claude_usd", "Claude", "USD",
         lambda x, p: (f"={i('claude_marginal_usd_per_order')}+IF({i('claude_count_list_price')}"
                       f"=1,{i('claude_list_usd_per_order')},0)"),
         "Marginal $0 on the subscription; list price only if the flag is 1"),
        ("sam3_usd", "SAM 3", "USD", lambda x, p: f"={i('sam3_usd_per_order')}", "Runs locally"),
        ("_s_material", "Material", None, None, ""),
        ("runs_slice", "Slice runs used (Counts = 1)", "runs",
         lambda x, p: (f"=COUNTIFS({_td('product')},{x}${PER_KEYS_ROW},{_td('counts')},1,"
                       f"{_td('tier')},\"m\",{_td('total_g')},\"<>\")"), "Count in Test data"),
        ("g_pla_bambu", "Bambu PLA grams (incl. purge)", "g",
         lambda x, p: _avg_f(x, "pla_bambu_g", "total_g",
                             f"{i(f'{p}.model_g')}+{i(f'{p}.purge_g')}"),
         "Average of slices, else model + purge placeholders"),
        ("g_pla_generic", "Generic PLA grams (incl. purge)", "g",
         lambda x, p: _avg_f(x, "pla_generic_g", "total_g", "0"), "Average of slices"),
        ("g_petg", "PETG grams (incl. purge)", "g",
         lambda x, p: _avg_f(x, "petg_g", "total_g", "0"), "Average of slices"),
        ("g_other", "Other grams (priced as Bambu PLA)", "g",
         lambda x, p: _avg_f(x, "other_g", "total_g", "0"), "Average of slices"),
        ("g_purge", "Purge grams (flush and prime tower)", "g",
         lambda x, p: _avg_f(x, "purge_g", "purge_g", i(f"{p}.purge_g")),
         "Average of slices, else the placeholder"),
        ("g_total", "Grams printed", "g",
         lambda x, p: (f"={L('g_pla_bambu', x)}+{L('g_pla_generic', x)}+{L('g_petg', x)}"
                       f"+{L('g_other', x)}+IF({i('slice_grams_include_purge')}=1,0,"
                       f"{L('g_purge', x)})"), "Sum of filaments"),
        ("g_model", "Model grams (without purge)", "g",
         lambda x, p: f"={L('g_total', x)}-{L('g_purge', x)}", "Printed minus purge"),
        ("units_per_plate", "Units per plate", "units",
         lambda x, p: f"=IF({i(f'{p}.units_per_plate')}=0,1,{i(f'{p}.units_per_plate')})", ""),
        ("g_purge_charged", "Purge grams charged per unit", "g",
         lambda x, p: f"={L('g_purge', x)}*{i('purge_scale')}/{L('units_per_plate', x)}",
         "Purge x purge scale / units per plate"),
        ("g_charged", "Grams charged per unit", "g",
         lambda x, p: f"={L('g_model', x)}+{L('g_purge_charged', x)}", ""),
        ("usd_per_g", "Blended filament price", "USD/g",
         lambda x, p: (f"=IF({L('g_total', x)}=0,0,({L('g_pla_bambu', x)}*{R('pla_bambu_g')}"
                       f"+{L('g_pla_generic', x)}*{R('pla_generic_g')}+{L('g_petg', x)}"
                       f"*{R('petg_g')}+{L('g_other', x)}*{R('pla_bambu_g')}"
                       f"+IF({i('slice_grams_include_purge')}=1,0,{L('g_purge', x)}"
                       f"*{R('pla_bambu_g')}))/{L('g_total', x)})"),
         "Weighted by grams of each filament"),
        ("material_usd", "Material", "USD",
         lambda x, p: f"={L('g_charged', x)}*{L('usd_per_g', x)}*{waste}*{fail}",
         "Grams x price x (1 + waste) x (1 + failed prints)"),
        ("_s_machine", "Printer time", None, None, ""),
        ("print_min", "Print minutes", "min",
         lambda x, p: _avg_f(x, "print_min", "print_min", i(f"{p}.print_min")),
         "Average of slices, else the placeholder"),
        ("power_usd", "Electricity", "USD",
         lambda x, p: f"={L('print_min', x)}/60*{R('power_h')}*{fail}", ""),
        ("depreciation_usd", "Depreciation and upkeep", "USD",
         lambda x, p: f"={L('print_min', x)}/60*{R('dep_h')}*{fail}", ""),
        ("machine_usd", "Machine", "USD",
         lambda x, p: f"={L('power_usd', x)}+{L('depreciation_usd', x)}", ""),
        ("_s_hardware", "Hardware", None, None, ""),
        ("hardware_usd", "Key ring, magnets, hook", "USD",
         lambda x, p: (f"={i(f'{p}.keyrings')}*{i('keyring_usd')}+{i(f'{p}.magnets')}"
                       f"*{i('magnet_usd')}+{i(f'{p}.hooks')}*{i('ornament_hook_usd')}"),
         "Count x unit price"),
        ("_s_labor", "Labor", None, None, ""),
        ("review_min", "Review", "min", lambda x, p: f"={i(f'{p}.review_min')}", ""),
        ("supports_min", "Remove supports, clean up", "min",
         lambda x, p: f"={i(f'{p}.supports_min')}", ""),
        ("packing_min", "Pack, attach hardware, label", "min",
         lambda x, p: f"={i(f'{p}.packing_min')}", ""),
        ("labor_min", "Labor minutes", "min",
         lambda x, p: f"={L('review_min', x)}+{L('supports_min', x)}+{L('packing_min', x)}", ""),
        ("labor_usd", "Labor", "USD",
         lambda x, p: f"={L('labor_min', x)}/60*{R('labor_h')}", "Minutes / 60 x hourly rate"),
        ("_s_total", "Production cost", None, None, ""),
        ("production_usd", "Production cost (before packaging, shipping, fees)", "USD",
         lambda x, p: (f"={L('meshy_usd', x)}+{L('claude_usd', x)}+{L('sam3_usd', x)}"
                       f"+{L('material_usd', x)}+{L('machine_usd', x)}+{L('hardware_usd', x)}"
                       f"+{L('labor_usd', x)}"), "Sum of the lines above"),
        ("_s_ship", "Packaging and shipping", None, None, ""),
        ("pack_ship_usd", "Packaging, shipped", "USD",
         lambda x, p: (f"=IF({i(f'{p}.uses_box')}=2,{i('packaging_box_large_usd')},"
                       f"IF({i(f'{p}.uses_box')}=1,{i('packaging_box_usd')},"
                       f"{i('packaging_mailer_usd')}))+{i('packaging_fill_usd')}"),
         "Mailer, small box or larger box + fill"),
        ("ship_oz", "Packed weight", "oz", lambda x, p: f"={i(f'{p}.ship_oz')}", ""),
        ("label_usd", "USPS Ground Advantage label", "USD",
         lambda x, p: (f"=IF({L('ship_oz', x)}<=4,{i('usps_ga_4oz_usd')},IF({L('ship_oz', x)}"
                       f"<=8,{i('usps_ga_8oz_usd')},{i('usps_ga_16oz_usd')}))"),
         "Tier by packed weight"),
        ("pack_pickup_usd", "Packaging, pickup or booth", "USD",
         lambda x, p: f"={i('packaging_pickup_usd')}", ""),
        ("_s_market", "Market (delivered price, Use = 1)", None, None, ""),
        ("market_n", "Listings, all regions", "count",
         lambda x, p: f"=COUNT({mk(p)})", "Large figure uses figurine listings"),
        ("market_q1", "Low quartile", "USD",
         lambda x, p: f"=IFERROR(QUARTILE({mk(p)},1),0)", ""),
        ("market_median", "Median", "USD", lambda x, p: f"=IFERROR(MEDIAN({mk(p)}),0)", ""),
        ("market_q3", "High quartile", "USD",
         lambda x, p: f"=IFERROR(QUARTILE({mk(p)},3),0)", ""),
        ("ny_n", "Listings, NY shops", "count", lambda x, p: f"=COUNT({mk(p, True)})", ""),
        ("ny_median", "Median, NY shops", "USD",
         lambda x, p: f"=IFERROR(MEDIAN({mk(p, True)}),0)", ""),
        ("ny_q3", "High quartile, NY shops", "USD",
         lambda x, p: f"=IFERROR(QUARTILE({mk(p, True)},3),0)", ""),
        ("_s_options", "Price options, own site", None, None, ""),
        ("own_break_even", "Cost floor (break-even price)", "USD",
         lambda x, p: f"={P}!$I${own_row(p)}", "From Prices"),
        ("own_target", "Cost plus target margin", "USD",
         lambda x, p: f"={P}!$J${own_row(p)}", "From Prices, before rounding"),
        ("own_market_median", "Market median", "USD",
         lambda x, p: f"={L('market_median', x)}", ""),
        ("own_price", "Own-site price", "USD", lambda x, p: f"={P}!$O${own_row(p)}",
         "From Prices: formula price or your override"),
        ("own_fees", "Card fees at that price", "USD", lambda x, p: f"={P}!$Q${own_row(p)}",
         "Stripe on price plus tax"),
        ("own_total_cost", "Total cost per order, own site", "USD",
         lambda x, p: f"={P}!$F${own_row(p)}+{P}!$Q${own_row(p)}",
         "Production + packaging + label + fees"),
        ("own_flag", "Own-site price vs NY high quartile", "",
         lambda x, p: (f'=IF(AND({L("ny_q3", x)}>0,{L("own_price", x)}>{L("ny_q3", x)}),'
                       f'"above NY high quartile","")'), "Flag"),
        ("own_margin_at_median", "Margin if priced at the market median", "share",
         lambda x, p: (f"=IF({L('market_median', x)}=0,0,({L('market_median', x)}"
                       f"-{P}!$F${own_row(p)}-({L('market_median', x)}*{P}!$G${own_row(p)}"
                       f"+{P}!$H${own_row(p)}))/{L('market_median', x)})"),
         "Negative = below break-even at that price"),
        ("own_margin_at_ny_median", "Margin if priced at the NY median", "share",
         lambda x, p: (f"=IF({L('ny_median', x)}=0,0,({L('ny_median', x)}"
                       f"-{P}!$F${own_row(p)}-({L('ny_median', x)}*{P}!$G${own_row(p)}"
                       f"+{P}!$H${own_row(p)}))/{L('ny_median', x)})"),
         "Negative = below break-even at that price"),
        ("_s_half", "Sensitivity: purge cut by half (material only)", None, None, ""),
        ("material_half_usd", "Material, purge halved", "USD",
         lambda x, p: (f"=({L('g_model', x)}+{L('g_purge', x)}*{i('purge_scale')}*0.5/"
                       f"{L('units_per_plate', x)})*{L('usd_per_g', x)}*{waste}*{fail}"), ""),
        ("production_half_usd", "Production cost, purge halved", "USD",
         lambda x, p: (f"={L('production_usd', x)}-{L('material_usd', x)}"
                       f"+{L('material_half_usd', x)}"), ""),
        ("own_price_half", "Own-site formula price, purge halved", "USD",
         lambda x, p: "=" + _round_f(
             f"MAX(({L('production_half_usd', x)}+IF({P}!$C${ch_own}=1,{L('pack_ship_usd', x)}"
             f"+{L('label_usd', x)},{L('pack_pickup_usd', x)})+{P}!$E${ch_own})/(1-{P}!$D$"
             f"{ch_own}-{i('target_margin')}),{L('market_q1', x)})", f"{P}!$G${ch_own}"),
         "Same rule as Prices"),
        ("own_price_half_change", "Change vs current formula price", "USD",
         lambda x, p: f"={L('own_price_half', x)}-{P}!$M${own_row(p)}", ""),
    ]


def rate_formula(lay: Layout, k: str) -> str:
    i = lambda key: _in(lay, key)
    return {
        "usd_credit": (f"=IF({i('meshy_use_topup_rate')}=1,{i('meshy_topup_usd_per_credit')},"
                       f"IF({i('meshy_plan_credits_month')}=0,0,{i('meshy_plan_usd_month')}/"
                       f"{i('meshy_plan_credits_month')}))"),
        "pla_bambu_g": f"={i('filament_bambu_pla_usd_kg')}/1000",
        "pla_generic_g": f"={i('filament_generic_pla_usd_kg')}/1000",
        "petg_g": f"={i('filament_petg_usd_kg')}/1000",
        "power_h": f"={i('printer_power_w')}/1000*{i('electricity_usd_kwh')}",
        "dep_h": (f"=IF({i('printer_life_hours')}=0,0,{i('printer_price_usd')}/"
                  f"{i('printer_life_hours')})+{i('printer_upkeep_usd_hour')}"),
        "labor_h": (f"=IF(ISNUMBER({i('labor_rate_jackson_usd_hour')}),"
                    f"{i('labor_rate_jackson_usd_hour')},{i('labor_rate_default_usd_hour')})"),
        "tax": f"={i('buyer_tax_pct')}",
    }[k]


def plan_layout(inputs: dict[str, dict]) -> Layout:
    rate_rows = {k: 5 + n for n, (k, _, _, _) in enumerate(RATES)}
    lay = Layout(inputs=inputs, per_rows={}, rate_rows=rate_rows, price_rows={}, channel_rows={},
                 market_day_rows={})
    lay.channel_rows = {c: 6 + n for n, c in enumerate(CHANNELS)}
    r = 6 + len(CHANNELS) + 3  # one blank row, the section label, the header
    for p in PRODUCTS:
        for c in CHANNELS:
            lay.price_rows[(p, c)] = r
            r += 1
    r += 3
    for p in PRODUCTS:
        lay.market_day_rows[p] = r
        r += 1
    r = PER_KEYS_ROW + 1
    for line in per_lines(lay):
        lay.per_rows[line[0]] = r
        r += 1
    r = SIZE_FIRST
    for p in SIZE_PRODUCTS:
        for t in TIERS:
            lay.size_rows[(p, t)] = r
            r += 1
    r += 3
    for k in ("min", "max", "basis"):
        lay.other_rows[k] = r
        r += 1
    return lay


def _unit_fmt(unit: str) -> str:
    if unit == "USD":
        return USD
    if unit in ("USD/g", "USD/credit"):
        return "$0.0000"
    if unit == "USD/hour":
        return USD3
    if unit == "share":
        return PCT
    if unit in ("g", "min", "oz", "units"):
        return NUM1
    if unit in ("runs", "count", "credits"):
        return NUM0
    return "General"


def build_per_product(ws: Worksheet, lay: Layout) -> None:
    _title(ws, "Per product",
           "Live formulas. Rebuilt by refresh; edit Inputs, Test data Use, or Market instead. "
           "Averages use Test data runs with Counts = 1; with none, the Inputs placeholders.")
    ws["A4"] = "Rates"
    ws["A4"].font = F_BOLD
    ws["A4"].fill = FILL_SECTION
    for n, (k, label, unit, how) in enumerate(RATES):
        r = 5 + n
        ws.cell(row=r, column=1, value=label).font = F_BASE
        ws.cell(row=r, column=2, value=unit).font = F_BASE
        c = ws.cell(row=r, column=3, value=rate_formula(lay, k))
        c.font = F_LINK
        c.number_format = {"USD/credit": "$0.0000", "USD/g": "$0.00000",
                           "share": "0.000%"}.get(unit, _unit_fmt(unit))
        ws.cell(row=r, column=4, value=how).font = F_NOTE
    hdr = PER_KEYS_ROW - 1
    _header(ws, hdr, ["Line", "Unit"] + [PRODUCT_LABELS[p] for p in PRODUCTS] + ["How"],
            [44, 9] + [14] * len(PRODUCTS) + [50])
    for p in PRODUCTS:
        c = ws[f"{per_col(p)}{PER_KEYS_ROW}"]
        c.value = p
        c.font = F_KEY
    ws.cell(row=PER_KEYS_ROW, column=1, value="Product key").font = F_KEY
    ws.freeze_panes = ws.cell(row=PER_KEYS_ROW + 1, column=3)
    for lid, label, unit, f, how in per_lines(lay):
        r = lay.per_rows[lid]
        a = ws.cell(row=r, column=1, value=label)
        if f is None:
            for j in range(1, len(PRODUCTS) + 4):
                ws.cell(row=r, column=j).fill = FILL_SECTION
            a.font = F_BOLD
            continue
        a.font = F_BASE
        ws.cell(row=r, column=2, value=unit).font = F_BASE
        for p in PRODUCTS:
            x = per_col(p)
            c = ws[f"{x}{r}"]
            c.value = f(x, p)
            c.font = F_BASE
            c.number_format = _unit_fmt(unit)
            c.border = BOX
            if lid in ("production_usd", "own_break_even", "own_target", "own_price",
                       "own_total_cost"):
                c.fill = FILL_RESULT
                c.font = F_BOLD
        ws.cell(row=r, column=len(PRODUCTS) + 3, value=how).font = F_NOTE


# ------------------------------------------------------------------ Prices

PRICE_COLS = [
    ("product", "Product", 24), ("channel", "Channel", 32), ("production", "Production cost", 11),
    ("pack_ship", "Packaging and shipping", 11), ("extra", "Booth and staffing per unit", 11),
    ("cost", "Cost before fees", 11), ("pct", "Percent fees", 9), ("fixed", "Fixed fees", 9),
    ("break_even", "Break-even price (cost floor)", 11),
    ("target", "Target-margin price", 11), ("market_q1", "Market low quartile", 11),
    ("market_median", "Market median", 11), ("formula_price", "Formula price", 11),
    ("override", "Override (yours)", 11), ("price", "Price", 11), ("basis", "Basis", 9),
    ("fees", "Fees at price", 10), ("profit", "Profit per unit", 10), ("margin", "Margin", 9),
    ("buyer_pays", "Buyer pays with tax", 11), ("key", "Row key", 22),
]
PR_LETTER = {k: get_column_letter(j) for j, (k, _, _) in enumerate(PRICE_COLS, start=1)}
CH_COLS = ["Channel key", "Channel", "Ships", "Percent fees (on price, incl. fee on tax)",
           "Fixed fee per order", "Booth and staffing per unit", "Price ending", "How"]


def channel_formulas(lay: Layout, c: str) -> tuple[str, str, str, str, str]:
    i = lambda k: _in(lay, k)
    tax = f"(1+{i('buyer_tax_pct')})"
    labor = f"{q(PER)}!$C${lay.rate_rows['labor_h']}"
    pct, fixed, extra, how = {
        "own_site": (f"={i('stripe_pct')}*{tax}", f"={i('stripe_fixed_usd')}", "=0",
                     "Stripe on price plus tax; free USPS shipping built in"),
        "own_pickup": (f"={i('stripe_pct')}*{tax}", f"={i('stripe_fixed_usd')}", "=0",
                       "Stripe on price plus tax; buyer picks up in NYC"),
        "etsy": ((f"={i('etsy_transaction_pct')}+{i('etsy_payment_pct')}*{tax}"
                  f"+{i('etsy_offsite_ads_pct')}*{i('etsy_offsite_ads_share')}"),
                 f"={i('etsy_listing_usd')}+{i('etsy_payment_fixed_usd')}", "=0",
                 ("Transaction on price; payment on price plus tax; Offsite Ads x share; listing "
                  "renewal + payment fixed")),
        "facebook": (f"={i('facebook_pickup_pct')}", "=0", "=0",
                     "Local pickup paid in person: no fee"),
        "ebay": (f"={i('ebay_fvf_pct')}*{tax}", f"={i('ebay_fixed_usd')}", "=0",
                 "Final value fee on price plus tax, plus per-order fee"),
        "market": (f"={i('square_pct')}*{tax}*{i('market_card_share')}",
                   f"={i('square_fixed_usd')}*{i('market_card_share')}",
                   (f"=IF({i('market_units_per_day')}=0,0,({i('market_booth_usd_day')}"
                    f"+{i('market_hours_per_day')}*{labor})/{i('market_units_per_day')})"),
                   "Square on card sales; booth fee and staffing spread over units per day"),
    }[c]
    return f"={i(f'{c}.ships')}", pct, fixed, extra, how


def build_prices(ws: Worksheet, lay: Layout, overrides: dict[tuple[str, str], object]) -> None:
    _title(ws, "Prices",
           "Price = the higher of the target-margin price and the market low quartile, rounded "
           "up to the channel's price ending, unless you type an Override (yellow). Refresh keeps "
           "overrides.")
    ws["A4"] = "Channel fees"
    ws["A4"].font = F_BOLD
    ws["A4"].fill = FILL_SECTION
    for j, lab in enumerate(CH_COLS, start=1):
        c = ws.cell(row=5, column=j, value=lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for c in CHANNELS:
        r = lay.channel_rows[c]
        ships, pct, fixed, extra, how = channel_formulas(lay, c)
        vals = [c, CHANNEL_LABELS[c], ships, pct, fixed, extra,
                f"={_in(lay, f'{c}.price_ending')}", how]
        fmts = ["General", "General", "0", PCT, USD, USD, "0.00", "General"]
        for j, (v, fm) in enumerate(zip(vals, fmts, strict=True), start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = F_KEY if j == 1 else (F_NOTE if j == 8 else F_BASE)
            cell.number_format = fm
            cell.border = BOX
    top = min(lay.price_rows.values())
    ws.cell(row=top - 2, column=1, value="Prices").font = F_BOLD
    ws.cell(row=top - 2, column=1).fill = FILL_SECTION
    for j, (_, lab, w) in enumerate(PRICE_COLS, start=1):
        c = ws.cell(row=top - 1, column=j, value=lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[top - 1].height = 42
    ws.freeze_panes = ws.cell(row=top, column=3)
    X = PR_LETTER
    margin = _in(lay, "target_margin")
    tax = _in(lay, "buyer_tax_pct")
    for (p, c), r in lay.price_rows.items():
        pc = per_col(p)
        cr = lay.channel_rows[c]
        per = lambda lid: f"{q(PER)}!${pc}${lay.per_rows[lid]}"  # noqa: B023
        row = {
            "product": PRODUCT_LABELS[p], "channel": CHANNEL_LABELS[c],
            "production": f"={per('production_usd')}",
            "pack_ship": (f"=IF($C${cr}=1,{per('pack_ship_usd')}+{per('label_usd')},"
                          f"{per('pack_pickup_usd')})"),
            "extra": f"=$F${cr}",
            "cost": f"={X['production']}{r}+{X['pack_ship']}{r}+{X['extra']}{r}",
            "pct": f"=$D${cr}", "fixed": f"=$E${cr}",
            "break_even": f"=({X['cost']}{r}+{X['fixed']}{r})/(1-{X['pct']}{r})",
            "target": f"=({X['cost']}{r}+{X['fixed']}{r})/(1-{X['pct']}{r}-{margin})",
            "market_q1": f"={per('market_q1')}", "market_median": f"={per('market_median')}",
            "formula_price": "=" + _round_f(f"MAX({X['target']}{r},{X['market_q1']}{r})",
                                            f"$G${cr}"),
            "override": overrides.get((p, c)),
            "price": (f"=IF(ISNUMBER({X['override']}{r}),{X['override']}{r},"
                      f"{X['formula_price']}{r})"),
            "basis": (f'=IF(ISNUMBER({X["override"]}{r}),"override",IF({X["target"]}{r}>='
                      f'{X["market_q1"]}{r},"cost","market"))'),
            "fees": f"={X['price']}{r}*{X['pct']}{r}+{X['fixed']}{r}",
            "profit": f"={X['price']}{r}-{X['cost']}{r}-{X['fees']}{r}",
            "margin": f"=IF({X['price']}{r}=0,0,{X['profit']}{r}/{X['price']}{r})",
            "buyer_pays": f"={X['price']}{r}*(1+{tax})",
            "key": f"{p}|{c}",
        }
        for j, (k, _, _) in enumerate(PRICE_COLS, start=1):
            cell = ws.cell(row=r, column=j, value=row[k])
            cell.border = BOX
            cell.font = F_BASE
            if k in ("pct", "margin"):
                cell.number_format = PCT
            elif k not in ("product", "channel", "basis", "key"):
                cell.number_format = USD
            if k in ("production", "market_q1", "market_median"):
                cell.font = F_LINK
            if k == "override":
                cell.fill = FILL_YELLOW
                cell.font = F_INPUT
            if k in ("price", "margin"):
                cell.font = F_BOLD
                cell.fill = FILL_RESULT
            if k == "key":
                cell.font = F_KEY
    # Market day break-even
    first = min(lay.market_day_rows.values())
    ws.cell(row=first - 2, column=1, value="Market day: units to cover the booth fee and "
                                           "staffing").font = F_BOLD
    ws.cell(row=first - 2, column=1).fill = FILL_SECTION
    hdrs = ["Product", "Market price", "Contribution per unit before booth and staffing",
            "Booth fee plus staffing per day", "Units per day to break even"]
    for j, lab in enumerate(hdrs, start=1):
        c = ws.cell(row=first - 1, column=j, value=lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[first - 1].height = 42
    i = lambda k: _in(lay, k)
    labor = f"{q(PER)}!$C${lay.rate_rows['labor_h']}"
    for p, r in lay.market_day_rows.items():
        m = lay.price_rows[(p, "market")]
        vals = [PRODUCT_LABELS[p], f"=${X['price']}${m}",
                f"=${X['price']}${m}-${X['production']}${m}-${X['pack_ship']}${m}-${X['fees']}${m}",
                f"={i('market_booth_usd_day')}+{i('market_hours_per_day')}*{labor}",
                f"=IF(C{r}>0,D{r}/C{r},0)"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_BASE
            c.border = BOX
            c.number_format = NUM1 if j == 5 else USD


def read_overrides(ws: Worksheet) -> dict[tuple[str, str], object]:
    out = {}
    kc = [k for k, _, _ in PRICE_COLS].index("key") + 1
    oc = [k for k, _, _ in PRICE_COLS].index("override") + 1
    for r in range(1, ws.max_row + 1):
        key = ws.cell(row=r, column=kc).value
        if isinstance(key, str) and "|" in key:
            v = ws.cell(row=r, column=oc).value
            if v not in (None, ""):
                p, c = key.split("|", 1)
                out[(p, c)] = v
    return out


# ------------------------------------------------------------------ Sizes

SIZE_HDR = 5
SIZE_FIRST = 6
SIZE_COLS = [
    ("product", "Product", 20), ("tier", "Tier", 6), ("mm", "Size mm", 8),
    ("scale", "Size / M", 8), ("runs", "Slice runs at this tier", 9), ("basis", "Basis", 22),
    ("model_g", "Model g", 9), ("purge_g", "Purge g", 9), ("print_min", "Print min", 9),
    ("material", "Material", 10), ("machine", "Machine", 10), ("production", "Production cost", 11),
    ("cost", "Cost before fees, own site", 11), ("break_even", "Break-even price", 11),
    ("target", "Target-margin price", 11), ("formula_price", "Formula price", 11),
    ("override", "Override (yours)", 11), ("price", "Price", 11), ("margin", "Margin", 9),
    ("key", "Row key", 16),
]
SZ = {k: get_column_letter(j) for j, (k, _, _) in enumerate(SIZE_COLS, start=1)}
OTHER_BASIS = ("Custom (Other) range: the cheapest XS own-site price, up to the L figurine "
               "price plus the complexity premium in Inputs.")


def size_formulas(lay: Layout, p: str, t: str) -> dict[str, object]:
    """Formulas for one product and tier. M reads Per product and Prices directly; other tiers
    use their own slices when Test data has them, else scale the M slice by size (ASSUMED)."""
    i = lambda k: _in(lay, k)
    r = lay.size_rows[(p, t)]
    rm = lay.size_rows[(p, M_TIER)]
    x = per_col(p)
    per = lambda lid: f"{q(PER)}!${x}${lay.per_rows[lid]}"
    o = lay.price_rows[(p, "own_site")]
    P = q(PRICES)
    ending = f"{P}!$G${lay.channel_rows['own_site']}"
    fail = f"(1+{i('failed_print_pct')})"
    waste = f"(1+{i('material_waste_pct')})"
    Z = lambda k: f"{SZ[k]}{r}"
    f: dict[str, object] = {
        "product": PRODUCT_LABELS[p], "tier": t.upper(), "mm": SIZE_TIERS_MM[p][t],
        "scale": f"={SZ['mm']}{r}/${SZ['mm']}${rm}", "key": f"{p}|{t}",
    }
    if t == M_TIER:
        f.update({
            "runs": f"={per('runs_slice')}",
            "basis": f'=IF({Z("runs")}>0,"slice","placeholder (ASSUMED)")',
            "model_g": f"={per('g_model')}", "purge_g": f"={per('g_purge')}",
            "print_min": f"={per('print_min')}",
            "formula_price": f"={P}!$O${o}", "override": None,
            "price": f"={Z('formula_price')}",
        })
    else:
        crit = (f"{_td('product')},\"{p}\",{_td('tier')},\"{t}\",{_td('counts')},1,"
                f"{_td('total_g')},\"<>\"")
        n = Z("runs")
        f.update({
            "runs": f"=COUNTIFS({crit})",
            "basis": f'=IF({n}>0,"slice","scaled from M (ASSUMED)")',
            "model_g": (f"=IF({n}>0,SUMIFS({_td('total_g')},{crit})/{n}-IF("
                        f"{i('slice_grams_include_purge')}=1,SUMIFS({_td('purge_g')},{crit})/{n},0),"
                        f"{per('g_model')}*{Z('scale')}^{i('size_model_g_exponent')})"),
            "purge_g": (f"=IF({n}>0,SUMIFS({_td('purge_g')},{crit})/{n},"
                        f"{per('g_purge')}*{Z('scale')}^{i('size_purge_exponent')})"),
            "print_min": (f"=IF({n}>0,SUMIFS({_td('print_min')},{crit})/{n},"
                          f"{per('print_min')}*{Z('scale')}^{i('size_time_exponent')})"),
            "formula_price": "=" + _round_f(f"{Z('target')}*{P}!$O${o}/{P}!$J${o}", ending),
            "price": f"=IF(ISNUMBER({Z('override')}),{Z('override')},{Z('formula_price')})",
        })
    f.update({
        "material": (f"=({Z('model_g')}+{Z('purge_g')}*{i('purge_scale')}/"
                     f"{per('units_per_plate')})*{per('usd_per_g')}*{waste}*{fail}"),
        "machine": (f"={Z('print_min')}/60*({q(PER)}!$C${lay.rate_rows['power_h']}"
                    f"+{q(PER)}!$C${lay.rate_rows['dep_h']})*{fail}"),
        "production": (f"={per('production_usd')}-{per('material_usd')}-{per('machine_usd')}"
                       f"+{Z('material')}+{Z('machine')}"),
        "cost": f"={Z('production')}+{P}!$D${o}+{P}!$E${o}",
        "break_even": f"=({Z('cost')}+{P}!$H${o})/(1-{P}!$G${o})",
        "target": f"=({Z('cost')}+{P}!$H${o})/(1-{P}!$G${o}-{i('target_margin')})",
        "margin": (f"=IF({Z('price')}=0,0,({Z('price')}-{Z('cost')}-({Z('price')}*{P}!$G${o}"
                   f"+{P}!$H${o}))/{Z('price')})"),
    })
    return f


def build_sizes(ws: Worksheet, lay: Layout, overrides: dict[tuple[str, str], object]) -> None:
    _title(ws, "Sizes",
           "Own-site price per size tier. M is the Prices sheet's own-site price (override M "
           "there). Other tiers use their own slices when Test data has them, else the M slice "
           "scaled by size with the exponents in Inputs (ASSUMED). Price = target-margin price "
           "x (M price / M target), rounded; your Override (yellow) wins. Refresh keeps overrides.")
    _header(ws, SIZE_HDR, [h for _, h, _ in SIZE_COLS], [w for _, _, w in SIZE_COLS])
    ws.row_dimensions[SIZE_HDR].height = 42
    ws.freeze_panes = ws.cell(row=SIZE_FIRST, column=3)
    for (p, t), r in lay.size_rows.items():
        f = size_formulas(lay, p, t)
        if t != M_TIER:
            f["override"] = overrides.get((p, t))
        for j, (k, _, _) in enumerate(SIZE_COLS, start=1):
            c = ws.cell(row=r, column=j, value=f.get(k))
            c.font = F_BASE
            c.border = BOX
            if k in ("model_g", "purge_g", "print_min"):
                c.number_format = NUM1
            elif k == "scale":
                c.number_format = "0.00"
            elif k == "margin":
                c.number_format = PCT
                c.font = F_BOLD
            elif k in ("material", "machine", "production", "cost", "break_even", "target",
                       "formula_price", "override", "price"):
                c.number_format = USD
            if k == "mm":
                c.font = F_INPUT
            if k == "override":
                if t == M_TIER:
                    c.value = "see Prices"
                    c.font = F_KEY
                else:
                    c.fill = FILL_YELLOW
                    c.font = F_INPUT
            if k == "price":
                c.font = F_BOLD
                c.fill = FILL_RESULT
            if k == "key":
                c.font = F_KEY
    ws.cell(row=lay.other_rows["min"] - 2, column=1, value="Custom (Other) orders").font = F_BOLD
    ws.cell(row=lay.other_rows["min"] - 2, column=1).fill = FILL_SECTION
    xs = [f"{SZ['price']}{lay.size_rows[(p, 'xs')]}" for p in SIZE_PRODUCTS]
    lfig = f"{SZ['price']}{lay.size_rows[('figurine', 'l')]}"
    ending = f"{q(PRICES)}!$G${lay.channel_rows['own_site']}"
    rows = {
        "min": ("Low end: cheapest XS price", f"=MIN({','.join(xs)})"),
        "max": ("High end: L figurine plus complexity premium",
                "=" + _round_f(f"{lfig}*(1+{_in(lay, 'other_complexity_premium')})", ending)),
        "basis": ("Basis", OTHER_BASIS),
    }
    for k, (label, v) in rows.items():
        r = lay.other_rows[k]
        ws.cell(row=r, column=1, value=label).font = F_BASE
        c = ws.cell(row=r, column=3, value=v)
        c.font = F_BOLD if k != "basis" else F_NOTE
        if k != "basis":
            c.number_format = USD
            c.fill = FILL_RESULT


def read_size_overrides(ws: Worksheet) -> dict[tuple[str, str], object]:
    kc = [k for k, _, _ in SIZE_COLS].index("key") + 1
    oc = [k for k, _, _ in SIZE_COLS].index("override") + 1
    out = {}
    for r in range(SIZE_FIRST, ws.max_row + 1):
        key = ws.cell(row=r, column=kc).value
        if isinstance(key, str) and "|" in key:
            v = ws.cell(row=r, column=oc).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                p, t = key.split("|", 1)
                out[(p, t)] = v
    return out


# ------------------------------------------------------------------ Changelog


def build_changelog(ws: Worksheet) -> None:
    _title(ws, "Changelog", "One row per change to prices or the model. Refresh appends a row "
                            "when any price moves.")
    _header(ws, HDR, ["Date", "What changed", "Why"], [20, 70, 70])


def append_changelog(ws: Worksheet, rows: list[tuple[str, str, str]]) -> None:
    r = max(ws.max_row, HDR) + 1
    for date, what, why in rows:
        for j, v in enumerate((date, what, why), start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_BASE
            c.alignment = WRAP
            c.border = BOX
        r += 1


def read_changelog(ws: Worksheet) -> list[tuple]:
    return [tuple(ws.cell(row=r, column=j).value for j in (1, 2, 3))
            for r in range(FIRST, ws.max_row + 1) if ws.cell(row=r, column=1).value]


# ------------------------------------------------------------------ whole workbook


def new_workbook(inputs: list[Input], listings: list[dict]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_ORDER:
        wb.create_sheet(name)
    build_inputs(wb[INPUTS], inputs)
    build_market(wb[MARKET], listings)
    build_changelog(wb[LOG])
    return wb


def replace_sheet(wb: Workbook, name: str) -> Worksheet:
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
    else:  # a sheet added in a later version goes where SHEET_ORDER puts it
        later = [n for n in SHEET_ORDER[SHEET_ORDER.index(name) + 1:] if n in wb.sheetnames]
        idx = wb.sheetnames.index(later[0]) if later else len(wb.sheetnames)
    if name in wb.sheetnames:
        wb.remove(wb[name])
    return wb.create_sheet(name, idx)


def note(ws: Worksheet, cell: str, text: str) -> None:
    ws[cell].comment = Comment(text, "refresh")
