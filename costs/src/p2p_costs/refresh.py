"""Refresh the cost model: uv run python costs/refresh.py

1. Open costs/cost_model.xlsx (or build it from defaults and research the first time).
2. Keep Inputs, Market and the Prices override column as they are (only append new input keys
   and new research listings).
3. Re-ingest test outputs into Test data (keeping each run's Use value).
4. Rebuild Per product and Prices with live formulas.
5. Compute the same numbers in Python, evaluate the formulas with LibreOffice, and stop if any
   number differs by more than half a cent.
6. Write costs/pricing.json and append a Changelog row when a price moves.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from . import lo
from .defaults import (
    CHANNEL_LABELS,
    CHANNELS,
    INPUTS,
    PRODUCT_LABELS,
    PRODUCTS,
    SIZE_PRODUCTS,
    SIZE_TIERS_MM,
    TIERS,
    WEB_CHANNEL,
    WEB_PRODUCTS,
)
from .ingest import (
    DEFAULT_EXPONENTS,
    RunRow,
    ingest_all,
    load_archive,
    merge_archive,
    save_archive,
)
from .model import Result, compute
from .workbook import (
    INPUTS as S_INPUTS,
)
from .workbook import (
    LOG,
    MARKET,
    PER,
    PRICE_COLS,
    PRICES,
    SIZE_COLS,
    SIZES,
    TEST,
    Layout,
    _td_layout_ok,
    append_changelog,
    append_market,
    append_missing_inputs,
    build_per_product,
    build_prices,
    build_sizes,
    effective_use,
    new_workbook,
    per_col,
    per_lines,
    plan_layout,
    read_inputs,
    read_market,
    read_overrides,
    read_size_overrides,
    read_test_uses,
    replace_sheet,
    update_test_data,
    write_market_formulas,
)

COSTS = Path(__file__).resolve().parents[2]  # costs/
TOL_USD = 0.005


class Mismatch(AssertionError):
    pass


def load_research(research_dir: Path) -> list[dict]:
    out: list[dict] = []
    for f in sorted(research_dir.glob("market_*.json")):
        try:
            data = json.loads(f.read_text())
        except (OSError, ValueError):
            continue
        rows = data.get("listings", []) if isinstance(data, dict) else data
        for d in rows:
            if isinstance(d, dict):
                out.append(d)
    return out


def _value_of(raw, cached):
    if isinstance(raw, str) and raw.startswith("="):
        return cached
    return raw


def _num(v) -> float | None:
    if isinstance(v, bool) or v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def edited_uses(sheet: dict[str, object], defaults: dict[str, object]) -> dict[str, object]:
    """Migration from the old layout, where refresh wrote a default into Use: keep only the values
    Jackson changed (the ones that differ from what refresh last wrote)."""
    out = {}
    for k, v in sheet.items():
        old = defaults.get(k)
        if old is None or _num(v) != _num(old):
            out[k] = v
    return out


# ------------------------------------------------------------------ check against formulas


def compare(path: Path, lay: Layout, res: Result) -> list[str]:
    wb = load_workbook(path, data_only=True)
    per, pr = wb[PER], wb[PRICES]
    bad: list[str] = []

    def check(where: str, got, want) -> None:
        if isinstance(want, str):
            if (got or "") != want:
                bad.append(f"{where}: workbook {got!r}, python {want!r}")
            return
        g = _num(got)
        if g is None or not math.isfinite(g) or abs(g - want) > max(TOL_USD, abs(want) * 1e-9):
            bad.append(f"{where}: workbook {got!r}, python {want:.6f}")

    for k, r in lay.rate_rows.items():
        check(f"{PER}!C{r} rate {k}", per[f"C{r}"].value, res.rates[k])
    for lid, _, _, f, _ in per_lines(lay):
        if f is None:
            continue
        r = lay.per_rows[lid]
        for p in PRODUCTS:
            want = res.per_product[p].get(lid)
            if want is None:
                bad.append(f"{PER} line {lid}: no python value")
                continue
            check(f"{PER}!{per_col(p)}{r} {lid}", per[f"{per_col(p)}{r}"].value, want)
    keys = [k for k, _, _ in PRICE_COLS]
    for (p, c), r in lay.price_rows.items():
        x = res.prices[(p, c)]
        for j, k in enumerate(keys, start=1):
            if k in ("product", "channel", "key", "override"):
                continue
            check(f"{PRICES}!{pr.cell(row=r, column=j).coordinate} {p}/{c} {k}",
                  pr.cell(row=r, column=j).value, x[k])
    for p, r in lay.market_day_rows.items():
        check(f"{PRICES}!E{r} market day {p}", pr[f"E{r}"].value, res.market_day[p]["units"])
    sz = wb[SIZES]
    skip = {"product", "tier", "key", "override", "mm"}
    for (p, t), r in lay.size_rows.items():
        z = res.sizes[(p, t)]
        for j, (k, _, _) in enumerate(SIZE_COLS, start=1):
            if k in skip:
                continue
            check(f"{SIZES}!{sz.cell(row=r, column=j).coordinate} {p}/{t} {k}",
                  sz.cell(row=r, column=j).value, z[k])
    for k in ("min", "max"):
        check(f"{SIZES}!C{lay.other_rows[k]} other {k}", sz[f"C{lay.other_rows[k]}"].value,
              res.other[k])
    return bad


# ------------------------------------------------------------------ changelog and outputs


def price_changes(prev: dict | None, res: Result) -> list[str]:
    old: dict[tuple[str, str], float] = {}
    if prev:
        for group in ("products", "reference"):
            for p, d in (prev.get(group) or {}).items():
                for c, v in (d.get("channels") or {}).items():
                    old[(p, c)] = v
                for t, v in (d.get("sizes") or {}).items():
                    old[(p, f"size {t}")] = v.get("price") if isinstance(v, dict) else None
        for k, v in (prev.get("other") or {}).items():
            if k in ("min", "max"):
                old[("other", k)] = v
    out = []
    for p in SIZE_PRODUCTS:
        for t in TIERS:
            if t == "m":
                continue
            new = round(float(res.sizes[(p, t)]["price"]), 2)
            was = old.get((p, f"size {t}"))
            if was is None or abs(was - new) >= 0.005:
                label = f"{PRODUCT_LABELS[p]}, size {t.upper()}"
                out.append(f"{label}: {'new' if was is None else f'${was:.2f}'} to ${new:.2f}")
    for k in ("min", "max"):
        new = round(float(res.other[k]), 2)
        was = old.get(("other", k))
        if was is None or abs(was - new) >= 0.005:
            out.append(f"Custom range {k}: {'new' if was is None else f'${was:.2f}'} to "
                       f"${new:.2f}")
    for p in PRODUCTS:
        for c in CHANNELS:
            new = round(float(res.prices[(p, c)]["price"]), 2)
            was = old.get((p, c))
            if was is None or abs(was - new) >= 0.005:
                label = f"{PRODUCT_LABELS[p]}, {CHANNEL_LABELS[c]}"
                out.append(f"{label}: {'new' if was is None else f'${was:.2f}'} to ${new:.2f}")
    return out


def reasons(prev_state: dict, inputs: dict, runs: dict, n_market: int, added: list[str],
            n_new_market: int, note: str | None) -> str:
    why = []
    if note:
        why.append(note)
    old_inputs = prev_state.get("inputs", {})
    changed = [k for k, v in inputs.items() if k in old_inputs and old_inputs[k] != v]
    if changed:
        why.append("inputs changed: " + ", ".join(changed[:12])
                   + (" and more" if len(changed) > 12 else ""))
    if added:
        why.append("new inputs: " + ", ".join(added[:8]))
    old_runs = prev_state.get("runs", {})
    new_runs = [k for k in runs if k not in old_runs]
    moved = [k for k in runs if k in old_runs and old_runs[k] != runs[k]]
    if new_runs:
        why.append(f"{len(new_runs)} new test runs")
    if moved:
        why.append(f"{len(moved)} test runs changed or now count differently")
    if n_new_market:
        why.append(f"{n_new_market} new market listings")
    elif prev_state and prev_state.get("market_rows") != n_market:
        why.append("market listings edited")
    return "; ".join(why) or "recomputed"


def pricing_payload(res: Result, now: str) -> dict:
    def entry(p: str) -> dict:
        x = res.prices[(p, WEB_CHANNEL)]
        d = {
            "price": round(float(x["price"]), 2),
            "from": True,
            "cost_floor": round(float(x["break_even"]), 2),
            "basis": x["basis"],
            "channels": {c: round(float(res.prices[(p, c)]["price"]), 2) for c in CHANNELS},
        }
        if p in SIZE_PRODUCTS:
            d["sizes"] = {t: {"price": round(float(res.sizes[(p, t)]["price"]), 2),
                              "mm": SIZE_TIERS_MM[p][t]} for t in TIERS}
        return d

    return {
        "updated_at": now,
        "currency": "USD",
        "products": {p: entry(p) for p in WEB_PRODUCTS},
        "reference": {p: entry(p) for p in PRODUCTS if p not in WEB_PRODUCTS},
        "other": {"min": round(float(res.other["min"]), 2),
                  "max": round(float(res.other["max"]), 2),
                  "basis": ("From the cheapest XS price to an L figurine plus a "
                            f"{res.other['premium']:.0%} premium for custom work.")},
        "notes": (
            "price is the own-site price with free USPS shipping; sales tax is added at checkout. "
            "from = true means a starting price (more subjects or a larger size cost more). "
            "cost_floor is the own-site break-even price. basis is cost (target margin), market "
            "(market low quartile) or override (set by hand). channels holds the price per "
            "channel: own_site, own_pickup (NYC pickup), etsy, facebook (local pickup), ebay, market "
            "(NYC booth). "
            "reference.large_figure is the Meshy-scale figure, not a listed product. Built by "
            "costs/refresh.py from costs/cost_model.xlsx."),
    }


# ------------------------------------------------------------------ main flow


def run(*, workbook: Path, repo_root: Path, extra_jobs: list[Path] | None = None,
        include_worktrees: bool = True, check: bool = True, note: str | None = None,
        now: datetime | None = None, ingest: bool = True) -> dict:
    costs = workbook.parent
    pricing_path = costs / "pricing.json"
    state_path = costs / "refresh_state.json"
    archive_path = costs / "test_data.jsonl"
    research = load_research(costs / "research")
    now = now or datetime.now(UTC)
    stamp = now.isoformat(timespec="seconds")

    existed = workbook.exists()
    if existed:
        wb = load_workbook(workbook)
        wbv = load_workbook(workbook, data_only=True)
    else:
        wb = new_workbook(INPUTS, research)
        wbv = None

    inputs = read_inputs(wb[S_INPUTS], wbv[S_INPUTS] if wbv else None)
    added = append_missing_inputs(wb[S_INPUTS], INPUTS, inputs) if existed else []

    n_new_market = append_market(wb[MARKET], read_market(wb[MARKET]), research) if existed else 0
    write_market_formulas(wb[MARKET])
    market = read_market(wb[MARKET])
    if wbv is not None:
        cached = {d["row"]: d for d in read_market(wbv[MARKET])}
        for d in market:
            c = cached.get(d["row"], {})
            for k in ("price_usd", "shipping_usd", "use", "region", "product"):
                d[k] = _value_of(d.get(k), c.get(k))

    overrides_raw = read_overrides(wb[PRICES]) if existed and PRICES in wb.sheetnames else {}
    overrides_val = (read_overrides(wbv[PRICES]) if wbv is not None and PRICES in wbv.sheetnames
                     else dict(overrides_raw))
    overrides_num = {k: float(v) for k, v in overrides_val.items() if _num(v) is not None}
    sizes_raw = read_size_overrides(wb[SIZES]) if SIZES in wb.sheetnames else {}
    sizes_val = (read_size_overrides(wbv[SIZES]) if wbv is not None and SIZES in wbv.sheetnames
                 else dict(sizes_raw))
    sizes_num = {k: float(v) for k, v in sizes_val.items() if _num(v) is not None}
    # Jackson's Use cells (values, so a formula typed there reads as its result).
    old_layout = existed and not _td_layout_ok(wb[TEST])
    uses = read_test_uses(wbv[TEST]) if wbv is not None and TEST in wbv.sheetnames else {}

    archive = load_archive(archive_path)
    migrated = edited_uses(uses, {k: r.auto for k, r in archive.items()}) if old_layout else {}
    if ingest:
        exps = {k: _num(inputs.get(f"size_{k}_exponent" if k != "model" else
                                   "size_model_g_exponent", {}).get("value")) or d
                for k, d in DEFAULT_EXPONENTS.items()}
        archive = merge_archive(archive, ingest_all(repo_root, extra_jobs, include_worktrees,
                                                    exps))
        save_archive(archive_path, archive)
    update_test_data(wb[TEST], list(archive.values()), migrated)
    use_now = migrated if old_layout else uses
    test_rows = []
    for r in archive.values():
        d = r.to_json()
        d["use"] = use_now.get(r.run_key)
        d["counts"] = effective_use(d["use"], r.auto)
        test_rows.append(d)

    lay = plan_layout(inputs)
    build_per_product(replace_sheet(wb, PER), lay)
    build_prices(replace_sheet(wb, PRICES), lay, overrides_raw)
    build_sizes(replace_sheet(wb, SIZES), lay, sizes_raw)

    values = {k: v["value"] for k, v in inputs.items()}
    res = compute(values, test_rows, market, overrides_num, sizes_num)

    prev = json.loads(pricing_path.read_text()) if pricing_path.exists() else None
    prev_state = json.loads(state_path.read_text()) if state_path.exists() else {}
    runs = {r["run_key"]: [r.get("counts"), r.get("credits"), r.get("total_g"), r.get("purge_g"),
                           r.get("print_min")] for r in test_rows}
    moves = price_changes(prev, res)
    log_rows = []
    if not existed:
        log_rows.append((stamp, "Workbook built. Prices set: " + "; ".join(moves),
                         note or "First build from defaults, research listings and test data"))
    elif moves:
        why = reasons(prev_state, values, runs, len(market), added, n_new_market, note)
        log_rows.append((stamp, "; ".join(moves), why))
    elif note:
        log_rows.append((stamp, "No price moved", note))
    append_changelog(wb[LOG], log_rows)

    tmp = costs / ".cost_model.tmp.xlsx"
    wb.save(tmp)
    try:
        if check:
            lo.recalc_in_place(tmp)
            bad = compare(tmp, lay, res)
            if bad:
                raise Mismatch("Python and workbook formulas disagree:\n  " + "\n  ".join(bad[:40]))
        tmp.replace(workbook)
    finally:
        if tmp.exists():
            tmp.unlink()

    pricing_path.write_text(json.dumps(pricing_payload(res, stamp), indent=2) + "\n")
    state_path.write_text(json.dumps({"inputs": values, "runs": runs,
                                      "market_rows": len(market)}, indent=1, default=str) + "\n")
    return {"result": res, "moves": moves, "added_inputs": added, "new_market": n_new_market,
            "runs": len(test_rows), "checked": check}


def summary(out: dict) -> str:
    res: Result = out["result"]
    lines = [f"{'product':<28}{'price':>9}{'floor':>9}{'margin':>8}  basis"]
    for p in PRODUCTS:
        x = res.prices[(p, WEB_CHANNEL)]
        lines.append(f"{PRODUCT_LABELS[p]:<28}{x['price']:>9.2f}{x['break_even']:>9.2f}"
                     f"{x['margin']:>8.0%}  {x['basis']}")
    lines.append(f"test runs: {out['runs']}; formulas checked: {out['checked']}; "
                 f"prices moved: {len(out['moves'])}")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Refresh costs/cost_model.xlsx and pricing.json")
    ap.add_argument("--workbook", type=Path, default=COSTS / "cost_model.xlsx")
    ap.add_argument("--root", type=Path, default=COSTS.parent, help="repo root to scan")
    ap.add_argument("--jobs-dir", type=Path, action="append", default=[],
                    help="extra jobs folder to ingest (repeatable)")
    ap.add_argument("--no-worktrees", action="store_true",
                    help="do not scan .claude/worktrees/* for test outputs")
    ap.add_argument("--no-check", action="store_true",
                    help="skip the LibreOffice formula check (not recommended)")
    ap.add_argument("--reason", help="why you are refreshing; goes in the Changelog")
    a = ap.parse_args(argv)
    try:
        out = run(workbook=a.workbook, repo_root=a.root, extra_jobs=a.jobs_dir,
                  include_worktrees=not a.no_worktrees, check=not a.no_check, note=a.reason)
    except Mismatch as e:
        print(e, file=sys.stderr)
        return 2
    print(summary(out))
    return 0


__all__ = ["RunRow", "main", "run"]
