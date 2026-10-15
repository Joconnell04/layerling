"""Refresh keeps Jackson's cells (Inputs, overrides, Use), re-ingests test outputs, and its Python
numbers match the workbook formulas as LibreOffice evaluates them."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from p2p_costs import lo
from p2p_costs.defaults import SIZE_TIERS_MM
from p2p_costs.refresh import compare, run
from p2p_costs.workbook import IC, INPUTS, PRICES, TD_COL, TEST, plan_layout, read_inputs

KEY_MM = SIZE_TIERS_MM["keychain"]  # tiers from contracts, never hard-coded here


def _tier_at(mm: float) -> str:
    return min(KEY_MM, key=lambda t: abs(KEY_MM[t] - mm))


def _row_for(ws, col: int, value: str) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == value:
            return r
    raise KeyError(value)


def _snapshot(ws) -> list[tuple]:
    return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                                 max_col=9)]


def _jackson_edits(path: Path) -> None:
    wb = load_workbook(path)
    ins = wb[INPUTS]
    ins.cell(row=_row_for(ins, IC["key"], "labor_rate_jackson_usd_hour"), column=IC["value"],
             value=30)
    ins.cell(row=_row_for(ins, IC["key"], "target_margin"), column=IC["notes"],
             value="Jackson's note")
    pr = wb[PRICES]
    pr.cell(row=_row_for(pr, 21, "keychain|own_site"), column=14, value=19.99)
    td = wb[TEST]
    # Force the flat Meshy keychain and the procedural lab keychain out, the old procedural
    # keychain slice in.
    td.cell(row=_row_for(td, 1, "slice:beagle45__z_combined__saving_infill_back"),
            column=TD_COL["use"], value=0)
    td.cell(row=_row_for(td, 1, "slice:keychain__z_combined__writer_defaults"),
            column=TD_COL["use"], value=0)
    td.cell(row=_row_for(td, 1, "slice:prod_keychain"), column=TD_COL["use"], value=1)
    wb.save(path)


def test_ingest_rows(fake_repo, costs_dir):
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    rows = {json.loads(x)["run_key"]: json.loads(x)
            for x in (costs_dir / "test_data.jsonl").read_text().splitlines()}
    k = rows["slice:prod_keychain"]
    assert k["product"] == "keychain" and k["auto"] == 0
    assert k["config"].startswith("settings not recorded")
    assert k["purge_g"] == 5.39 and k["print_min"] == pytest.approx(34.15)
    assert k["pla_generic_g"] == pytest.approx(2.01) and k["pla_bambu_g"] == pytest.approx(1.83)
    assert k["petg_g"] == pytest.approx(4.74)
    assert rows["slice:meshy_multicolor_x2d"]["product"] == "large_figure"
    assert rows["job:j1"]["credits"] == 66 and rows["job:j1"]["product"] == "nameplate"
    assert rows["job:j1"]["auto"] == 1 and rows["job:j2"]["auto"] == 0
    assert "production cost profile" in rows["job:j1"]["config"]
    ledger = [r for r in rows.values() if r["kind"] == "meshy spend"]
    assert len(ledger) == 1 and ledger[0]["credits"] == 18 and ledger[0]["auto"] == 0


def test_production_config_decides_auto(fake_repo, costs_dir):
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    rows = {json.loads(x)["run_key"]: json.loads(x)
            for x in (costs_dir / "test_data.jsonl").read_text().splitlines()}
    flat = rows["slice:beagle45__z_combined__saving_infill_back"]
    assert flat["product"] == "keychain" and flat["auto"] == 1
    assert flat["config"].startswith("production defaults") and "flat" in flat["config"]
    assert flat["tier"] == _tier_at(45) and flat["size_mm"] == 45
    base = rows["slice:beagle45__baseline__base"]
    assert base["auto"] == 0
    assert "prime_volume_mode=Default (production Saving)" in base["config"]
    assert "upright (production flat)" in base["config"]
    # The procedural keychain has no recorded size, so it sits at M. It counts there unless a
    # real Meshy slice lands on the same tier.
    proc = rows["slice:keychain__z_combined__writer_defaults"]
    assert proc["config"].startswith("production defaults") and proc["tier"] == "m"
    assert proc["auto"] == (0 if _tier_at(45) == "m" else 1)
    assert rows["slice:meshy_multicolor_x2d"]["auto"] == 0


def test_refresh_keeps_inputs_overrides_and_use(fake_repo, costs_dir, slice_report):
    wbp = costs_dir / "cost_model.xlsx"
    run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=False)
    _jackson_edits(wbp)
    wb = load_workbook(wbp)
    ins = wb[INPUTS]
    gone = _row_for(ins, IC["key"], "market_card_share")
    ins.delete_rows(gone)  # a key Jackson removed by accident comes back at the bottom
    wb.save(wbp)
    before = _snapshot(load_workbook(wbp)[INPUTS])

    sr = fake_repo / "data" / "slice_reports"
    (sr / "prod_magnet.slice_report.json").write_text(json.dumps(slice_report))
    out = run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=False,
              note="test refresh")

    wb = load_workbook(wbp)
    after = _snapshot(wb[INPUTS])
    assert after[:len(before)] == before, "Inputs rows changed"
    assert after[len(before)][IC["key"] - 1] == "market_card_share"
    assert out["added_inputs"] == ["market_card_share"]
    pr = wb[PRICES]
    assert pr.cell(row=_row_for(pr, 21, "keychain|own_site"), column=14).value == 19.99
    td = wb[TEST]
    use = TD_COL["use"]
    assert td.cell(row=_row_for(td, 1, "slice:beagle45__z_combined__saving_infill_back"),
                   column=use).value == 0
    assert td.cell(row=_row_for(td, 1, "slice:prod_keychain"), column=use).value == 1
    assert td.cell(row=_row_for(td, 1, "slice:beagle45__baseline__base"),
                   column=use).value is None, "refresh wrote into a blank Use cell"
    new = _row_for(td, 1, "slice:prod_magnet")
    assert td.cell(row=new, column=use).value is None and td.cell(row=new, column=1).value
    assert td.cell(row=new, column=TD_COL["counts"]).value.startswith("=IF(ISNUMBER(")
    pricing = json.loads((costs_dir / "pricing.json").read_text())
    kc = pricing["products"]["keychain"]
    assert kc["price"] == 19.99 and kc["basis"] == "override"
    assert set(pricing) >= {"updated_at", "currency", "products", "notes"}
    assert set(kc) >= {"price", "from", "cost_floor", "basis"}
    log = [r for r in wb["Changelog"].iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(log) == 2 and "test refresh" in log[-1][2]


@pytest.mark.skipif(not lo.available(), reason="LibreOffice not installed")
def test_python_matches_formulas(fake_repo, costs_dir):
    wbp = costs_dir / "cost_model.xlsx"
    run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=False)
    _jackson_edits(wbp)
    out = run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=True)
    assert out["checked"]
    res = out["result"]
    assert res.prices[("keychain", "own_site")]["price"] == 19.99
    assert res.rates["labor_h"] == 30
    # Use overrides: the flat beagle is forced out, the old procedural slice forced in.
    assert res.per_product["keychain"]["g_total"] == pytest.approx(8.58)
    assert res.per_product["keychain"]["runs_slice"] == 1
    # The saved workbook carries LibreOffice's values, and a wrong Python number is caught.
    lay = plan_layout(read_inputs(load_workbook(wbp)[INPUTS]))
    assert compare(wbp, lay, res) == []
    wrong = copy.deepcopy(res)
    wrong.prices[("magnet", "etsy")]["target"] += 1.0
    wrong.per_product["figurine"]["material_usd"] += 0.02
    bad = compare(wbp, lay, wrong)
    assert len(bad) == 2


def test_migration_keeps_only_edited_use_values():
    from p2p_costs.refresh import edited_uses

    sheet = {"a": 1, "b": 0, "c": "1", "d": 1}
    last_written = {"a": 1, "b": 1, "c": 1}  # b was edited by Jackson; d is unknown to refresh
    assert edited_uses(sheet, last_written) == {"b": 0, "d": 1}


def test_geometry_production_slice_supersedes_lab_runs(fake_repo, costs_dir, slice_report):
    # Same size as the lab beagle, so both land on the same contracts tier.
    cfg = {"source": "p2p_geometry round 2", "product": "keychain", "target_size_mm": 45.0,
           "orientation": "flat_face_up", "supports": "none", "layer_mm": 0.2,
           "process_overrides": {"prime_volume_mode": "Saving", "flush_into_infill": "1",
                                 "enable_support": "0"}}
    sr = fake_repo / "data" / "slice_reports"
    (sr / "geometry_keychain__baseline__production.json").write_text(
        json.dumps(dict(slice_report, config=cfg)))
    upright = dict(cfg, product="magnet", orientation="upright")  # magnets print flat
    (sr / "geometry_magnet__baseline__production.json").write_text(
        json.dumps(dict(slice_report, config=upright)))
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    rows = {json.loads(x)["run_key"]: json.loads(x)
            for x in (costs_dir / "test_data.jsonl").read_text().splitlines()}
    geo = rows["slice:geometry_keychain__baseline__production"]
    assert geo["auto"] == 1 and "geometry production slice" in geo["config"]
    lab = rows["slice:beagle45__z_combined__saving_infill_back"]
    assert geo["tier"] == lab["tier"] == _tier_at(45)
    assert lab["auto"] == 0 and "superseded by the geometry production slice" in lab["notes"]
    mag = rows["slice:geometry_magnet__baseline__production"]
    assert mag["auto"] == 0 and "upright (production flat)" in mag["config"]


def test_size_tiers_in_pricing_json_and_overrides(fake_repo, costs_dir, slice_report):
    from p2p_costs.workbook import SIZE_COLS, SIZES

    cfg = {"source": "p2p_geometry", "product": "keychain", "size_tier": "xs",
           "target_size_mm": KEY_MM["xs"], "orientation": "flat_face_up", "supports": "none",
           "layer_mm": 0.2, "process_overrides": {"prime_volume_mode": "Saving",
                                                  "flush_into_infill": "1"}}
    sr = fake_repo / "data" / "slice_reports"
    small = dict(slice_report, print_time_s=600.0,
                 filaments=[{"slot": 1, "filament_id": "generic-pla-black", "grams": 3.0}],
                 total_grams=3.0, purge_grams=1.0)
    (sr / "geometry_keychain_xs__production.json").write_text(json.dumps(dict(small, config=cfg)))
    wbp = costs_dir / "cost_model.xlsx"
    out = run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=False)
    res = out["result"]
    xs = res.sizes[("keychain", "xs")]
    assert xs["basis"] == "slice" and xs["print_min"] == pytest.approx(10.0)
    assert res.sizes[("keychain", "l")]["basis"] == "scaled from M (ASSUMED)"
    # The XS slice does not leak into the M averages (M has the procedural keychain only).
    assert res.per_product["keychain"]["print_min"] == pytest.approx(2049.0 / 60, abs=0.01)
    pricing = json.loads((costs_dir / "pricing.json").read_text())
    for p, d in pricing["products"].items():
        sizes = d["sizes"]
        assert list(sizes) == ["xs", "s", "m", "l"]
        assert sizes["m"]["price"] == d["price"], p
        prices = [sizes[t]["price"] for t in ("xs", "s", "m", "l")]
        if p != "keychain":  # scaled from M only; the fake keychain slices are unrelated shapes
            assert prices == sorted(prices), p
    other = pricing["other"]
    assert 0 < other["min"] < other["max"] and other["basis"]

    wb = load_workbook(wbp)
    sz = wb[SIZES]
    kc = [k for k, _, _ in SIZE_COLS].index("key") + 1
    oc = [k for k, _, _ in SIZE_COLS].index("override") + 1
    sz.cell(row=_row_for(sz, kc, "magnet|l"), column=oc, value=49.0)
    wb.save(wbp)
    out = run(workbook=wbp, repo_root=fake_repo, include_worktrees=False, check=False)
    sz = load_workbook(wbp)[SIZES]
    assert sz.cell(row=_row_for(sz, kc, "magnet|l"), column=oc).value == 49.0
    assert out["result"].sizes[("magnet", "l")]["price"] == 49.0
    pricing = json.loads((costs_dir / "pricing.json").read_text())
    assert pricing["products"]["magnet"]["sizes"]["l"]["price"] == 49.0


def test_slices_land_on_contracts_tiers(fake_repo, costs_dir, slice_report):
    """A slice made near a tier's size is scaled to it; one far from every tier does not count;
    a writer's tier label loses to the size it was actually made at."""
    base = {"source": "p2p_geometry", "product": "keychain", "orientation": "flat_face_up",
            "supports": "none", "layer_mm": 0.2,
            "process_overrides": {"prime_volume_mode": "Saving", "flush_into_infill": "1"}}
    sr = fake_repo / "data" / "slice_reports"
    near = KEY_MM["l"] * 1.05
    far = KEY_MM["xs"] * 0.7
    runs = {"geometry_keychain_l__production": dict(base, size_tier="l", target_size_mm=near),
            "geometry_keychain_xs__production": dict(base, size_tier="xs", target_size_mm=far),
            "geometry_keychain_s__production": dict(base, size_tier="m",
                                                    target_size_mm=KEY_MM["s"])}
    for name, cfg in runs.items():
        (sr / f"{name}.json").write_text(json.dumps(dict(slice_report, config=cfg)))
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    rows = {json.loads(x)["run_key"]: json.loads(x)
            for x in (costs_dir / "test_data.jsonl").read_text().splitlines()}
    lrow = rows["slice:geometry_keychain_l__production"]
    assert lrow["tier"] == "l" and lrow["auto"] == 1 and "scaled to L" in lrow["config"]
    assert lrow["print_min"] < 2049.0 / 60  # scaled down to the smaller tier size
    xrow = rows["slice:geometry_keychain_xs__production"]
    assert xrow["auto"] == 0 and "not a current keychain size" in xrow["config"]
    srow = rows["slice:geometry_keychain_s__production"]
    assert srow["tier"] == "s" and "labelled M" in srow["config"]


def test_removed_slice_file_stops_counting(fake_repo, costs_dir):
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    (fake_repo / "data" / "slice_reports" / "keychain__z_combined__writer_defaults.json").unlink()
    run(workbook=costs_dir / "cost_model.xlsx", repo_root=fake_repo, include_worktrees=False,
        check=False)
    rows = {json.loads(x)["run_key"]: json.loads(x)
            for x in (costs_dir / "test_data.jsonl").read_text().splitlines()}
    gone = rows["slice:keychain__z_combined__writer_defaults"]
    assert gone["auto"] == 0 and "removed" in gone["notes"]
    assert rows["job:j1"]["auto"] == 1  # job runs keep counting
